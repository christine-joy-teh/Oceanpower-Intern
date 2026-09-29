from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs" / "proposals"

SOURCES = [
    (
        DOCS / "sources" / "OCEANPOWER_AI_CHATBOT_DEMO_PROPOSAL_EN.md",
        DOCS / "supporting" / "Oceanpower_AI_Chatbot_Demo_Proposal_EN.docx",
        "Oceanpower Bilingual AI Chatbot Demo Proposal",
        "Management decision draft for an internal demonstration",
        "Prepared 24 September 2026",
        "en",
    ),
    (
        DOCS / "sources" / "OCEANPOWER_AI_CHATBOT_DEMO_PROPOSAL_CN.md",
        DOCS / "supporting" / "Oceanpower_AI_Chatbot_Demo_Proposal_CN.docx",
        "Oceanpower 中英双语 AI 聊天机器人演示方案",
        "供管理层内部演示决策使用",
        "编制日期 2026 年 9 月 24 日",
        "zh",
    ),
]

NAVY = "17365D"
PALE_BLUE = "EAF1F8"
PALE_GRAY = "F5F6F7"
BORDER = "D9D9D9"
TEXT = RGBColor(31, 31, 31)


def set_run_font(run, name, size=None, bold=None, color=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=110, start=120, bottom=110, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:color"), BORDER)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def prevent_row_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    cant_split.set(qn("w:val"), "true")
    tr_pr.append(cant_split)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def add_hyperlink(paragraph, text, url, font_name):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    new_run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), font_name)
    r_fonts.set(qn("w:hAnsi"), font_name)
    r_fonts.set(qn("w:eastAsia"), font_name)
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    r_pr.extend([r_fonts, color, underline])
    new_run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    new_run.append(text_node)
    hyperlink.append(new_run)
    paragraph._p.append(hyperlink)


def add_rich_text(paragraph, text, font_name, size=10.4):
    url_match = re.search(r"(https?://\S+)$", text)
    if url_match:
        prefix = text[: url_match.start()].rstrip()
        if prefix:
            run = paragraph.add_run(prefix + " ")
            set_run_font(run, font_name, size=size, color=TEXT)
        add_hyperlink(paragraph, url_match.group(1), url_match.group(1), font_name)
        return

    parts = re.split(r"(\*\*.*?\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, font_name, size=size, bold=True, color=TEXT)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, font_name, size=size, color=TEXT)


def configure_document(doc, font_name, language, footer_text=None):
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(1.9)
    section.left_margin = Cm(2.35)
    section.right_margin = Cm(2.35)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = font_name
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
    normal.font.size = Pt(10.4 if language == "en" else 10.5)
    normal.font.color.rgb = TEXT
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.18

    for style_name, size, before, after in (
        ("Title", 25, 0, 16),
        ("Heading 1", 16, 15, 7),
        ("Heading 2", 12.5, 11, 5),
    ):
        style = styles[style_name]
        style.font.name = font_name
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font_name)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        if style_name == "Title":
            p_pr = style._element.get_or_add_pPr()
            p_bdr = p_pr.find(qn("w:pBdr"))
            if p_bdr is not None:
                p_pr.remove(p_bdr)

    for section in doc.sections:
        footer = section.footer
        p = footer.paragraphs[0]
        default_footer = "Oceanpower AI Chatbot Demo Proposal" if language == "en" else "Oceanpower AI 聊天机器人演示方案"
        p.text = (footer_text or default_footer) + "    "
        for run in p.runs:
            set_run_font(run, font_name, size=8, color=RGBColor(100, 100, 100))
        add_page_number(p)


def add_cover(doc, title, subtitle, prepared, font_name):
    for _ in range(4):
        doc.add_paragraph()
    p = doc.add_paragraph(style="Title")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    set_run_font(run, font_name, size=25, bold=True, color=RGBColor(0, 0, 0))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(subtitle)
    set_run_font(run, font_name, size=13, color=RGBColor(70, 70, 70))
    p.paragraph_format.space_before = Pt(8)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(prepared)
    set_run_font(run, font_name, size=10.5, color=RGBColor(100, 100, 100))
    p.paragraph_format.space_before = Pt(18)

    doc.add_page_break()


def add_markdown_table(doc, rows, font_name):
    if not rows:
        return
    col_count = max(len(row) for row in rows)
    table = doc.add_table(rows=len(rows), cols=col_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_borders(table)
    available = 6.85
    if col_count == 3:
        widths = [available * 0.22, available * 0.34, available * 0.44]
    elif col_count == 2:
        widths = [available * 0.32, available * 0.68]
    else:
        widths = [available / col_count] * col_count

    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx]
        prevent_row_split(row)
        if r_idx == 0:
            set_repeat_table_header(row)
        for c_idx in range(col_count):
            cell = row.cells[c_idx]
            cell.width = Inches(widths[c_idx])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            if r_idx == 0:
                set_cell_shading(cell, NAVY)
            elif r_idx % 2 == 0:
                set_cell_shading(cell, PALE_BLUE)
            else:
                set_cell_shading(cell, "FFFFFF")
            text = row_data[c_idx] if c_idx < len(row_data) else ""
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.1
            run = p.add_run(text)
            set_run_font(
                run,
                font_name,
                size=8.8 if r_idx else 9,
                bold=r_idx == 0,
                color=RGBColor(255, 255, 255) if r_idx == 0 else TEXT,
            )
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def markdown_to_docx(source_path, output_path, title, subtitle, prepared, language):
    font_name = "Aptos" if language == "en" else "Microsoft YaHei"
    doc = Document()
    if "Programme" in title or "项目建议书" in title:
        footer_text = "Oceanpower AI Programme Proposal" if language == "en" else "Oceanpower AI 项目建议书"
    else:
        footer_text = None
    configure_document(doc, font_name, language, footer_text)
    doc.core_properties.title = title
    doc.core_properties.subject = subtitle
    doc.core_properties.author = "Oceanpower internship project"
    add_cover(doc, title, subtitle, prepared, font_name)

    lines = source_path.read_text(encoding="utf-8").splitlines()
    # Skip title and metadata already presented on the cover.
    start = next((i for i, line in enumerate(lines) if line.startswith("## ")), 0)
    lines = lines[start:]
    paragraph_buffer = []
    in_code = False
    code_lines = []
    index = 0

    def flush_paragraph():
        nonlocal paragraph_buffer
        if paragraph_buffer:
            p = doc.add_paragraph()
            add_rich_text(p, " ".join(paragraph_buffer), font_name)
            paragraph_buffer = []

    while index < len(lines):
        line = lines[index].rstrip()

        if line.startswith("```"):
            flush_paragraph()
            if not in_code:
                in_code = True
                code_lines = []
            else:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.line_spacing = 1.05
                run = p.add_run("\n".join(code_lines))
                set_run_font(run, "Consolas", size=9, color=TEXT)
                in_code = False
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if line.startswith("## "):
            flush_paragraph()
            doc.add_heading(line[3:].strip(), level=1)
            index += 1
            continue
        if line.startswith("### "):
            flush_paragraph()
            doc.add_heading(line[4:].strip(), level=2)
            index += 1
            continue

        if line.startswith("|"):
            flush_paragraph()
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                current = lines[index].strip()
                if not re.fullmatch(r"\|?[\s:\-|]+\|?", current):
                    cells = [cell.strip() for cell in current.strip("|").split("|")]
                    table_lines.append(cells)
                index += 1
            add_markdown_table(doc, table_lines, font_name)
            continue

        bullet_match = re.match(r"^-\s+(.*)$", line)
        numbered_match = re.match(r"^(\d+)\.\s+(.*)$", line)
        if bullet_match or numbered_match:
            flush_paragraph()
            p = doc.add_paragraph(style="List Bullet" if bullet_match else None)
            p.paragraph_format.space_after = Pt(3)
            if numbered_match:
                p.paragraph_format.left_indent = Inches(0.25)
                p.paragraph_format.first_line_indent = Inches(-0.25)
                text = f"{numbered_match.group(1)}. {numbered_match.group(2)}"
            else:
                text = bullet_match.group(1)
            add_rich_text(p, text, font_name)
            index += 1
            continue

        if not line.strip():
            flush_paragraph()
        else:
            paragraph_buffer.append(line.strip())
        index += 1

    flush_paragraph()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


if __name__ == "__main__":
    for source, output, title, subtitle, prepared, language in SOURCES:
        markdown_to_docx(source, output, title, subtitle, prepared, language)
        print(output)

